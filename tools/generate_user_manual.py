# -*- coding: utf-8 -*-
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "多平台采集工作台操作说明-v2.1.1.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn("w:" + key))
        if node is None:
            node = OxmlElement("w:" + key)
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_row_cant_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_row_repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_run_font(run, name="Microsoft YaHei", size=10.5, color="222222", bold=False):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold


def set_para(paragraph, before=0, after=6, line=1.18, keep=False):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if keep:
        fmt.keep_with_next = True


def add_text(doc, text, style=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    set_para(p, after=6)
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    set_para(p, after=3)
    r = p.add_run(text)
    set_run_font(r)
    return p


_number_counter = 0


def add_number(doc, text):
    """Add a locally numbered step so each procedure starts at 1."""
    global _number_counter
    _number_counter += 1
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(0.35)
    p.paragraph_format.first_line_indent = Cm(-0.35)
    set_para(p, after=3)
    r = p.add_run(f"{_number_counter}. ")
    set_run_font(r)
    r = p.add_run(text)
    set_run_font(r)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0].cells
    for i, value in enumerate(headers):
        header[i].text = ""
        p = header[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_para(p, after=0, line=1.0)
        r = p.add_run(str(value))
        set_run_font(r, size=9.5, color="FFFFFF", bold=True)
        set_cell_shading(header[i], "1F4E78")
        set_cell_border(header[i])
        set_cell_margins(header[i])
        header[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        set_row_cant_split(table.rows[-1])
        for i, value in enumerate(row):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 or len(headers) <= 2 else WD_ALIGN_PARAGRAPH.LEFT
            set_para(p, after=0, line=1.08)
            r = p.add_run(str(value))
            set_run_font(r, size=9.2)
            set_cell_shading(cells[i], "F3F7FB" if ridx % 2 else "FFFFFF")
            set_cell_border(cells[i])
            set_cell_margins(cells[i], top=105, bottom=105)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_row_repeat_header(table.rows[0])
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level=1):
    global _number_counter
    _number_counter = 0
    p = doc.add_paragraph(style=f"Heading {level}")
    set_para(p, before=10 if level == 1 else 6, after=5, keep=True)
    r = p.add_run(text)
    set_run_font(r, size=16 if level == 1 else 12.5, color="000000", bold=True)
    return p


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = paragraph.add_run("第 ")
    set_run_font(r, size=9, color="666666")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)
    r = paragraph.add_run(" 页")
    set_run_font(r, size=9, color="666666")


def build():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(1.65)
    sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(1.8)
    sec.right_margin = Cm(1.8)
    sec.header_distance = Cm(0.7)
    sec.footer_distance = Cm(0.7)

    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    for name, size in (("Heading 1", 16), ("Heading 2", 12.5), ("Heading 3", 11.5)):
        style = doc.styles[name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
    for style_name in ("List Bullet", "List Bullet 2", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(10.5)

    footer = sec.footer.paragraphs[0]
    add_page_number(footer)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, before=18, after=5, line=1.0)
    r = p.add_run("多平台采集工作台操作说明")
    set_run_font(r, size=24, color="000000", bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_para(p, after=14, line=1.0)
    r = p.add_run("适用于 2.1.1 及兼容版本")
    set_run_font(r, size=11, color="666666")

    add_text(doc, "本说明面向日常使用人员，介绍从登录、账号绑定、创建采集任务，到线索筛选、互动回复、内容发布、消息同步、诊断和更新的完整操作。员工只需要按界面提示执行，不需要修改源码或数据库。", bold_prefix="本说明面向日常使用人员")
    add_text(doc, "软件默认把真实发送和真实发布作为受控操作。采集、回复、私信和发布都应先使用模拟或预览流程确认目标页面、账号和内容，确认无误后再由有权限的人员开启真实动作。")

    add_heading(doc, "目录", 1)
    toc = [
        "1 使用前准备",
        "2 登录注册与个人中心",
        "3 界面结构与工作区切换",
        "4 账号管理",
        "5 关键词组与采集任务",
        "6 任务中心与采集过程",
        "7 采集总览",
        "8 线索中心",
        "9 互动中心",
        "10 发布工作区",
        "11 消息中心",
        "12 设置与诊断",
        "13 数据文件日志与更新",
        "14 推荐标准流程",
        "15 常见问题处理",
    ]
    for item in toc:
        add_bullet(doc, item)

    add_heading(doc, "1 使用前准备", 1)
    add_text(doc, "软件运行在 Windows 环境。免安装发布包可直接双击多平台采集工作台.exe；源码运行需要 Python 3.11，浏览器控制需要 websockets 依赖。采集、回复、私信和发布还需要 BitBrowser 本地 API 服务以及已安装的 Chrome。")
    add_table(doc, ["项目", "要求", "检查方式"], [
        ("软件", "使用管理员分发的免安装包，首次启动会自动创建 data 目录。", "确认 EXE 与 runtime 位于同一软件目录。"),
        ("BitBrowser", "本地 API 已启动，默认地址为 http://127.0.0.1:54345。", "在设置与诊断中点击检测端口。"),
        ("浏览器窗口", "每个平台先在 BitBrowser 创建窗口，记录窗口 ID，并完成登录。", "账号管理中能看到窗口并显示已打开或已绑定。"),
        ("平台账号", "账号绑定到正确的平台和窗口。", "账号管理中确认平台、昵称、窗口 ID。"),
        ("LLM", "只在需要智能分析或内容生成时配置，未配置时使用本地规则或模板。", "设置中检测 API 连通性。"),
    ], widths=[3.2, 9.0, 4.4])
    add_text(doc, "浏览器窗口必须尽量保持一个窗口对应一个账号。打开浏览器后如果看到导航页，应先确认账号管理中的窗口 ID、平台和登录状态，再进行任务。不要把一个已登录其他账号的窗口重新绑定给当前账号。")

    add_heading(doc, "2 登录注册与个人中心", 1)
    add_heading(doc, "2.1 登录", 2)
    add_number(doc, "输入账号和密码，员工账号还需要经过管理员审批。")
    add_number(doc, "需要长期使用时勾选记住登录状态。登录状态只保存会话信息，不保存明文密码。")
    add_number(doc, "登录成功后，右上角个人入口显示员工姓名。不同员工的数据按登录账号隔离。")
    add_text(doc, "如果登录后看到上一个员工的账号、任务或线索，先退出登录并重新进入，不要继续操作。管理员可以查看全体员工数据；普通员工只能查看自己的数据。")
    add_heading(doc, "2.2 注册与审批", 2)
    add_text(doc, "注册时填写登录账号、六位以上密码和员工姓名。提交后状态为待审批，管理员在个人中心的用户审批中批准或拒绝。员工收不到审批提示时，应让管理员在管理员中心刷新待审批列表，并确认管理员使用的是统一账号服务。")
    add_heading(doc, "2.3 个人中心", 2)
    add_table(doc, ["功能", "用途"], [
        ("个人信息", "查看当前登录员工姓名、角色和登录状态。"),
        ("数据上传", "按管理员部署的同步方案手动上传本机数据；默认不会自动上传。"),
        ("退出登录", "清除当前会话并切换员工；退出后必须重新登录才能读取该员工数据。"),
        ("用户审批", "管理员查看待审批员工并批准、拒绝或禁用账号。"),
        ("管理员中心", "管理员查看员工、设备、备份、审计记录，并执行备份校验或恢复。"),
    ], widths=[4.0, 12.6])

    add_heading(doc, "3 界面结构与工作区切换", 1)
    add_text(doc, "左上角有采集和发布两个工作区按钮。切换工作区只切换业务界面，不会自动启动采集、回复或发布。后台任务继续由服务运行，页面以局部更新方式显示变化。")
    add_table(doc, ["工作区", "主要页面", "用途"], [
        ("采集", "采集总览、任务中心、账号管理、线索中心、互动中心、设置与诊断", "管理关键词、采集评论、筛选线索并进行后续互动。"),
        ("发布", "账号信息、内容生成、发布、消息中心、设置", "读取账号自己的内容、生成或编辑文案、审核发布并同步互动消息。"),
    ], widths=[3.0, 6.0, 7.6])
    add_text(doc, "顶部的刷新按钮通常只刷新当前页面数据，不等同于重启任务。任务中心的暂停、继续、停止和删除按钮会改变后台状态，操作前要确认任务编号。")

    add_heading(doc, "4 账号管理", 1)
    add_text(doc, "账号管理负责维护平台账号、BitBrowser 窗口和二者的绑定关系。账号管理中的账号是后续采集、回复、私信和发布所使用的账号，不是线索评论中的用户。")
    add_heading(doc, "4.1 创建浏览器窗口", 2)
    add_number(doc, "在账号管理点击创建浏览器，填写窗口名称和平台信息，提交后等待列表刷新。")
    add_number(doc, "确认新窗口出现在列表中，记录窗口 ID。必要时点击打开浏览器，完成平台登录。")
    add_number(doc, "登录完成后点击刷新账号或读取昵称，确认昵称、平台和窗口 ID 对应正确。")
    add_heading(doc, "4.2 添加和绑定账号", 2)
    add_number(doc, "选择平台，选择已经创建的窗口，填写账号标识或昵称。贴吧账号使用 API 配置，不要求绑定浏览器窗口。")
    add_number(doc, "点击添加账号或绑定账号。成功后应看到已绑定窗口和账号状态。")
    add_number(doc, "抖音昵称从个人主页读取，小红书从创作服务平台读取，微博从个人主页读取，B站从头像菜单读取，快手从个人主页左下角读取。若昵称读取不准，先确认页面已登录且打开的是该账号主页，再刷新。")
    add_table(doc, ["状态", "含义", "处理"], [
        ("空闲", "账号未被任务占用。", "可以分配给新任务或执行测试。"),
        ("工作中", "账号正在采集、回复、私信或发布。", "等待当前动作完成，不要重复打开同一账号。"),
        ("冷却中", "账号完成一批操作，正在按冷却策略等待。", "等待冷却结束。回复默认间隔 5 秒，连续 10 条后冷却 1 分钟。"),
        ("待人工验证", "后台检测到登录、验证码或平台阻断，需要人工处理。", "到对应浏览器完成处理，再回任务中心点击继续。"),
        ("未绑定或未打开", "窗口存在但还没有和账号正确绑定，或窗口未启动。", "检查窗口 ID、平台和 BitBrowser 服务。"),
    ], widths=[3.5, 7.0, 6.1])
    add_text(doc, "任务暂停、结束、失败或人工接管后，账号应释放。若页面仍显示工作中，先刷新账号列表；不要直接删除正在运行的窗口。删除窗口前必须确认不是其他任务正在使用。")

    add_heading(doc, "5 关键词组与采集任务", 1)
    add_heading(doc, "5.1 关键词组", 2)
    add_text(doc, "关键词组是可重复使用的搜索词集合。核心词之间用顿号或分隔符填写，必要时补充同义词、地区词和排除词。关键词组只保存搜索规则，不会自动开始任务。")
    add_number(doc, "在新建任务中点击新建关键词组，填写名称、适用平台和关键词内容。")
    add_number(doc, "保存后回到新建任务，在关键词组下拉框选择刚保存的词组。若下拉框没有内容，先刷新关键词组或确认当前登录账号下确实保存过词组。")
    add_number(doc, "执行任务时，组内每个关键词依次搜索。目标数量是每个关键词的目标数量，不是整个词组共用一个数量。")
    add_heading(doc, "5.2 新建任务", 2)
    add_table(doc, ["字段", "填写方式"], [
        ("平台", "选择抖音、小红书、B站、微博、快手或百度贴吧。不同平台会显示对应的排序选项。"),
        ("搜索关键词或关键词组", "直接填写关键词，或选择已经保存的预制词组。使用词组时，任务会逐个搜索所有词。"),
        ("搜索排序", "按平台可用规则选择，例如综合、最新或热门。快手和贴吧的可选项较少，以界面实际显示为准。"),
        ("采集模式", "通常使用标准采集。页面加载和评论滚动由平台适配器处理。"),
        ("每关键词目标", "每个关键词希望采集的作品数量。实际数量受平台搜索结果和无更多内容影响。"),
        ("每批数量", "每轮详情处理的作品数量，用于控制速度和资源占用。"),
        ("冷却秒数", "同一账号连续操作之间的等待时间。"),
        ("监控间隔", "定时增量监控的间隔，创建任务后可在任务中心单独设置。"),
        ("采集内容", "选择作品基础信息、作者信息、评论、互动数据、用户信息、地区和评论者主页等字段。"),
        ("参与账号", "只勾选对应平台的已登录账号。贴吧使用已配置 API 的账号。"),
        ("输出目录", "可以自由选择导出目录。任务数据本身仍保存在软件 data 目录。"),
    ], widths=[4.2, 12.4])
    add_text(doc, "创建任务后，平台和账号不会自动替换为其他账号。多个关键词必须按顺序完成当前关键词的搜索采集，只有达到目标数量或明确检测到平台提示“暂时没有更多了”后，才进入下一个关键词。")

    add_heading(doc, "6 任务中心与采集过程", 1)
    add_heading(doc, "6.1 任务状态", 2)
    add_table(doc, ["界面状态", "含义", "按钮动作"], [
        ("待采集", "任务尚未开始或等待账号。", "开始采集。"),
        ("搜索采集中", "正在逐个关键词搜索作品。", "暂停采集。"),
        ("评论采集中", "搜索结果已进入作品详情，正在滚动读取评论。", "暂停采集。"),
        ("暂停", "用户暂停或意外退出后保留断点。", "开始或继续。"),
        ("待人工验证", "需要人工处理登录或验证码。", "人工处理后点击继续。"),
        ("暂无更多视频", "该关键词已到平台尽头，或平台明确返回无更多提示。", "继续采集可重新检查是否有新增内容；不会把该关键词误判为已达到目标。"),
        ("已完成", "所有关键词达到目标，或各关键词已经明确无更多内容并完成后续处理。", "查看、导出或进入线索中心。"),
        ("失败", "任务发生无法继续的异常。", "查看失败原因，修复账号或页面后重新开始。"),
    ], widths=[4.0, 8.0, 4.6])
    add_heading(doc, "6.2 搜索和评论采集规则", 2)
    add_text(doc, "搜索阶段会滚动页面并等待动态内容加载。页面出现新作品时立即处理，不会每次固定等待十秒；翻到页面底部后，如果没有新内容，最长等待十秒再次确认。平台明确显示“暂时没有更多了”时才记录无更多结果。补采时相同 URL 会去重，重复旧视频不能单独作为无更多的依据。")
    add_text(doc, "评论区会采集当前页面可以加载的全部评论，包括楼中楼回复。抖音 NOTE 类型页面需要先聚焦或打开评论区，再滚动评论容器；其他平台也会检测评论容器并执行翻页或滚动加载。文字匹配会忽略图片、表情和装饰内容；空评论、纯图片、纯表情和纯数字评论不会发送给 LLM。")
    add_heading(doc, "6.3 暂停和继续", 2)
    add_number(doc, "正常工作时点击暂停，按钮应变为开始或继续，账号释放，当前关键词和已发现作品保留。")
    add_number(doc, "人工验证时到对应浏览器处理登录或验证码。处理完成后回任务中心点击继续，不要点击新建任务。")
    add_number(doc, "如果软件意外退出，重新打开后后台会把残留工作状态恢复为暂停。点击继续会从未完成的关键词或评论断点继续，不会直接跳过原关键词。")
    add_heading(doc, "6.4 导出与统计", 2)
    add_text(doc, "点击导出任务会导出有效作品、评论、意向和用户聚合数据。空标题、空评论和无效数据不会进入有效导出统计，因此任务完成数量应与导出文件中的有效数量一致。任务卡片显示任务编号、关键词、平台、账号、有效作品进度、评论数、起止时间和状态。")

    add_heading(doc, "7 采集总览", 1)
    add_text(doc, "采集总览是实时驾驶舱，展示进行中任务、有效作品、评论入库数量和待处理任务。任务队列按固定列显示关键词、平台、账号、进度、评论数和状态。后台事件采用局部刷新，某一条任务变化时只更新对应卡片或行。")
    add_table(doc, ["区域", "查看内容"], [
        ("顶部统计", "进行中任务、已完成有效作品、已采集评论、等待开始或人工处理的任务。"),
        ("任务队列", "每个任务的关键词、平台、账号、进度条、评论数和状态。"),
        ("实时事件", "任务开始、暂停、人工介入、无更多内容、导出和异常等事件。"),
        ("采集趋势", "按时间查看评论采集量变化。"),
        ("平台分布", "查看各平台的评论量和接入状态；未接入平台显示为预留。"),
    ], widths=[4.2, 12.4])
    add_text(doc, "如果总览长时间显示评论数为零，先看实时事件和任务中心状态。评论会按批次写入数据库并推送界面，不代表页面没有采集；若超过一个批次仍无变化，再检查浏览器页面和日志。")

    add_heading(doc, "8 线索中心", 1)
    add_text(doc, "线索中心保存采集到的评论、昵称、平台、地区、评论时间、原作地址和意向结果。列表中可以按任务、平台、地区、时效、意向和关键词并行筛选，筛选不会修改原始数据。")
    add_heading(doc, "8.1 筛选", 2)
    add_number(doc, "选择任务，可只查看某个采集任务产生的评论。")
    add_number(doc, "选择平台、地区、时效或意向等级。")
    add_number(doc, "在关键词筛选中输入昵称、评论或词组相关文字。关键词筛选与意向等级独立生效。")
    add_number(doc, "点击刷新按钮后读取最新入库数据；切换筛选项不会触发浏览器采集。")
    add_heading(doc, "8.2 选择和处理", 2)
    add_table(doc, ["操作", "结果"], [
        ("选择一条或多条评论", "右侧线索详情显示昵称、平台、地区、意向、评论原文和评论时间。"),
        ("点击评论时间表头", "按时间升序或降序排列。"),
        ("点击查看原作", "调用本机 Chrome 打开原作地址，并尽量定位到对应评论；如果平台需要滚动，会继续加载评论区。"),
        ("加入互动中心", "后台加入所选线索，当前线索页面保持不变；加入后不会从线索中心删除。"),
        ("加入发私信", "后台为所选用户创建私信草稿，当前线索页面保持不变。"),
        ("导出选中或按任务导出", "生成包含任务编号、线索信息、意向等级、平台、评论原文、地址和时间的文件。"),
    ], widths=[5.2, 11.4])
    add_text(doc, "手动选择加入互动中心或发私信时，不受自动意向和地区限制。因为这是人工筛选后的明确动作，后续仍然会要求账号、目标页面和输入框定位成功。")

    add_heading(doc, "9 互动中心", 1)
    add_text(doc, "互动中心处理两类动作：评论回复和发私信。每条记录保留来源平台图标、任务编号、用户昵称、评论原文、回复或私信完整正文、账号、状态和失败原因。")
    add_heading(doc, "9.1 状态和基本流程", 2)
    add_table(doc, ["状态", "操作"], [
        ("待生成", "进入发送页生成或编辑完整话术，也可以删除。"),
        ("待发送", "选择账号、模拟发送或真实发送，也可以删除；支持多选后一键处理。"),
        ("已回复", "查看已成功记录，按平台和账号筛选；保留原评论和回复内容。"),
        ("失败", "查看失败原因，可退回待生成或删除。"),
    ], widths=[4.0, 12.6])
    add_heading(doc, "9.2 模板和内容", 2)
    add_number(doc, "在模板管理中新增模板，填写模板名称、完整回复正文和变量。变量使用中文标签，可自定义添加可用变量。")
    add_number(doc, "在待生成或待发送记录中选择模板，模板完整内容会进入编辑框，不显示简述。")
    add_number(doc, "人工修改编辑框后先保存，再提交审核。批准时保存的正文才是最终发送正文。")
    add_number(doc, "评论回复和发私信可以互相切换。每条记录提供转到私信或转到评论回复入口。")
    add_heading(doc, "9.3 模拟与真实发送", 2)
    add_text(doc, "真实发送开关默认关闭。关闭时点击模拟或一键回复，程序打开目标平台、定位来源作品或用户、滚动加载评论、点击回复或私信按钮并填入内容，但不点击最终发送。开启真实发送后，只有审核通过并进入待发送的记录才会点击最终发送按钮。")
    add_text(doc, "回复和私信默认每条间隔五秒，连续十条后冷却一分钟。发送前会重新确认当前页面、目标用户、原评论和输入框；窗口位置或大小变化时，优先使用页面元素定位，坐标点击只作为第二道防御。")
    add_text(doc, "如果失败，必须查看记录中的阶段和原因，例如目标作品未找到、评论未加载、私信按钮未找到、输入框未找到、发送按钮未找到、账号为私密账号或平台页面未登录。不要直接重复点击真实发送。")

    add_heading(doc, "10 发布工作区", 1)
    add_heading(doc, "10.1 账号信息", 2)
    add_text(doc, "账号信息只读取当前账号自己发布的作品、图文内容、点赞数和评论详情，不读取采集任务中的他人作品。选择平台和账号后点击同步账号作品，程序会从账号主页或创作后台重新读取，页面残留的搜索结果不会作为账号作品。")
    add_text(doc, "不同平台的账号主页入口不同：抖音使用已登录账号的个人主页，小红书使用创作服务平台，B站从登录账号主页读取，微博进入个人主页，快手从个人主页读取。同步前确认账号已经登录。")
    add_heading(doc, "10.2 内容生成", 2)
    add_number(doc, "在内容生成页面填写选题、受众、平台和写作要求。")
    add_number(doc, "配置 LLM 时生成内容会调用智能 API；未配置时按本地模板生成。")
    add_number(doc, "生成结果中的思维链或 think 内容只用于排查，不会进入正式标题、正文或回复框。保存前检查正式正文。")
    add_number(doc, "点击加入发布或从内容生成导入，把完整标题、正文、话题和素材带入发布中心。")
    add_heading(doc, "10.3 发布编辑和审核", 2)
    add_table(doc, ["步骤", "说明"], [
        ("新建内容", "创建发布草稿，选择一个或多个平台版本。"),
        ("编辑平台版本", "分别编辑标题、完整正文、话题标签和素材。不同平台可以使用不同版本内容。"),
        ("保存平台版本", "保存当前版本；发布时会再次比对本地正文与浏览器输入内容。"),
        ("提交审核", "草稿进入待审核。"),
        ("批准", "审核后的正文会作为最终版本保存。"),
        ("进入待发布", "把批准版本放入发布队列。"),
    ], widths=[4.2, 12.4])
    add_heading(doc, "10.4 定时发布和真实发布", 2)
    add_number(doc, "在发布设置选择账号和发布时间。定时日期时间使用北京时间，日历不能选择过去时间。")
    add_number(doc, "保存定时发布后，发布队列会显示计划时间、账号和状态。")
    add_number(doc, "真实发布开关默认关闭。关闭时只进行预览或填充测试；开启后才会执行最终发布按钮。")
    add_number(doc, "平台素材类型必须匹配：视频素材进入视频模式，图片素材进入图文模式。发布按钮需要等待平台检测完成，程序会先等待冷却，再尝试点击并检查成功提示。")
    add_text(doc, "若发布后网页已经填入内容但未发布，先检查平台版本编辑中的正文是否保存、真实发布是否开启、账号是否正确，以及发布页是否出现检测等待。不要在检测期间连续点击发布。")

    add_heading(doc, "11 消息中心", 1)
    add_text(doc, "消息中心从平台自己的消息入口全量读取可见消息，并按同一用户归类为会话。左侧显示昵称或用户 ID，右侧显示详情。消息类型包括评论、回复、点赞、@提及、关注、私信、群通知和系统通知，具体类别随平台页面实际可见内容变化。")
    add_number(doc, "在发布工作区进入消息中心，选择平台和账号。")
    add_number(doc, "点击同步消息，等待各平台读取完成。同步失败项会单独显示，先查看失败平台和原因。")
    add_number(doc, "点击左侧会话查看右侧详情。已发布作品的评论互动可以继续转入互动中心处理。")
    add_text(doc, "消息同步是只读操作，不会自动回复、点赞、删除或发送私信。真实回复仍需进入互动中心并经过人工确认。")

    add_heading(doc, "12 设置与诊断", 1)
    add_heading(doc, "12.1 BitBrowser 连接", 2)
    add_number(doc, "填写 BitBrowser 地址和端口，默认地址通常为 http://127.0.0.1:54345。")
    add_number(doc, "点击保存连接设置，再点击检测并识别端口。只有检测成功后，账号管理和真实浏览器操作才有可靠基础。")
    add_number(doc, "如果端口连接失败，确认 BitBrowser 已启动、端口未被其他服务占用，并检查地址是否为本机地址。")
    add_heading(doc, "12.2 智能 API", 2)
    add_text(doc, "智能 API 用于评论意向分析、内容生成和回复草稿。填写服务商、API 地址、模型名称和 API Key 后保存，再点击 API 连通性检测。检测成功只代表接口可连通，不代表每种批量返回格式都正确。")
    add_text(doc, "批量意向分析按每批 500 条评论发送。程序要求模型返回只包含编号和意向等级的 JSON 数组，并覆盖每一条输入；模型返回 think、Markdown 代码块或解释文字时，程序会先尝试清理和解析，仍无法解析则记录原始片段供排查。未配置 LLM 时使用本地规则，空评论、图片、表情包和纯数字不发送给 LLM。")
    add_heading(doc, "12.3 健康诊断", 2)
    add_table(doc, ["按钮", "作用"], [
        ("运行平台检查", "检查已接入平台的基础状态，包括贴吧 API 配置状态。"),
        ("检测端口", "检查 BitBrowser API 是否可连接。"),
        ("真实浏览器诊断", "进入每个平台入口，确认页面、登录、评论区和关键控件；不发送真实内容。"),
        ("一键深度诊断", "使用已有或手动输入的一条评论，测试定位、填入和页面反馈，只填不发。"),
        ("导出诊断包", "导出脱敏后的状态、日志和截图，便于在其他电脑排查。"),
        ("日志统计", "按正常、警告、错误查看当天日志和下一条警告。"),
        ("检查更新", "从线上清单读取版本和构建标识。发现更新后再点击立即更新并重启。"),
    ], widths=[4.4, 12.2])
    add_heading(doc, "12.4 员工数据同步", 2)
    add_text(doc, "同步服务器地址、设备名称和服务器令牌属于部署参数，未配置时可以留空。默认同步不会自动上传，只有启用同步并点击立即同步才会执行。同步内容按登录员工隔离，通常包括员工自己的账号绑定、任务、作品、评论、线索、互动和发布记录，不上传密码、Cookie、API Key 等敏感信息。具体服务地址由管理员统一提供。")
    add_heading(doc, "12.5 运行日志", 2)
    add_text(doc, "日志记录任务、账号、浏览器、搜索、滚动、定位、点击、输入、发送、发布、API 调用摘要、返回阶段和异常反馈。敏感信息按规则脱敏。日志每小时归档保存到本地，设置页支持当天查看、统计和导出。")

    add_heading(doc, "13 数据文件日志与更新", 1)
    add_table(doc, ["位置", "内容"], [
        ("data\\platform_gui.db", "本机账号、任务、评论、线索、互动、发布和配置数据。"),
        ("data\\logs", "后台操作日志、小时归档和浏览器诊断记录。"),
        ("data\\exports", "任务、线索、互动和日志导出文件。"),
        ("logs\\update.log", "更新器下载、校验、替换和重启记录。"),
        ("code_backups", "代码更新前自动保存的程序文件，不包含业务数据。"),
    ], widths=[5.4, 11.2])
    add_text(doc, "升级只替换程序代码和运行资源，data、账号、任务、LLM 配置和日志目录应保留。正式升级前关闭软件和浏览器操作窗口，升级后重新打开，程序会自动读取原有数据。")
    add_text(doc, "免安装包缺少更新器时，把便携更新程序文件夹放到软件目录内，双击更新程序.bat。该脚本会把 update.ps1 放到多平台采集工作台.exe 同级目录，然后执行更新。旧版程序只识别根目录 update.ps1，不要只把脚本留在嵌套文件夹里。")

    add_heading(doc, "14 推荐标准流程", 1)
    add_number(doc, "启动 BitBrowser，确认 API 服务正常。")
    add_number(doc, "打开软件并登录员工账号，确认右上角姓名和数据范围正确。")
    add_number(doc, "在账号管理创建或刷新对应平台窗口，登录平台并绑定账号。")
    add_number(doc, "在新建任务中选择平台、关键词组、排序、目标数量、采集字段和参与账号。")
    add_number(doc, "创建任务后先运行小数量测试，确认搜索页面、作品详情和评论区能正常读取。")
    add_number(doc, "在任务中心观察搜索采集、评论采集、评论数、账号和状态；遇到人工验证就处理后点击继续。")
    add_number(doc, "在采集总览或线索中心确认评论入库，使用任务、平台、地区、意向和关键词筛选。")
    add_number(doc, "人工选择有效线索加入互动中心或发私信，确认原评论、昵称和原作地址。")
    add_number(doc, "选择模板并审核完整话术，先关闭真实发送做模拟填入测试。")
    add_number(doc, "确认页面定位、输入框、按钮和内容无误后，由有权限人员开启真实发送或真实发布。")
    add_number(doc, "完成后查看已回复、失败原因、消息中心和导出数据。")

    add_heading(doc, "15 常见问题处理", 1)
    add_table(doc, ["现象", "优先检查", "处理方式"], [
        ("打开的是导航页", "BitBrowser 窗口 ID、平台和登录态。", "在账号管理打开正确窗口，确认平台主页，再重新绑定；不要把导航页当作目标页面。"),
        ("提示需要人工验证但网页没有验证码", "是否处于抖音 NOTE 或其他特殊页面，评论区是否未展开。", "先进入正确作品并点击评论入口，再重新诊断；不要只依据滚轮无反应判定验证码。"),
        ("搜索只返回少量作品就进入详情", "任务是否提前判定无更多、页面是否还在加载。", "查看实时日志；只有达到目标或明确出现无更多提示才结束关键词。"),
        ("暂停后继续跳到下一个关键词", "当前关键词状态和任务断点。", "检查日志中的关键词状态；未达到目标且未记录无更多时，应继续当前关键词。"),
        ("线索加入互动中心后页面没有数据", "是否点击了加入互动中心，当前登录员工是否一致。", "点击后等待后台成功日志并刷新互动中心；员工切换后数据不会互相显示。"),
        ("互动中心找不到原评论", "原作地址、文字匹配、评论是否需要滚动加载。", "关闭真实发送，重新模拟；程序会忽略图片和表情并滚动加载评论。"),
        ("私信点成消息中心", "目标主页上的私信按钮与右上角全局消息按钮。", "确认目标用户主页，再执行模拟私信；右上角消息不是目标私信入口。"),
        ("内容已填入但没有发送", "真实发送开关、平台检测冷却、最终按钮是否出现。", "模拟模式本来就不点击发送；真实模式等待平台检测完成，再查看日志中的发送按钮定位。"),
        ("批量 LLM 分析无法解析", "原始返回是否含 think、Markdown 或非 JSON。", "查看日志片段；要求模型只返回编号和意向两个字段的 JSON 数组，或暂时使用本地规则。"),
        ("导出目录为空", "任务是否完成、输出目录权限、导出是否成功。", "点击导出后查看日志和 data\\exports；无效空标题和空评论会被过滤。"),
        ("切换员工仍看到原数据", "是否真正退出旧账号、登录会话是否恢复。", "退出登录后重新登录，确认右上角姓名；必要时关闭并重启软件。"),
        ("点击检查更新提示缺少更新器", "软件根目录是否有 update.ps1。", "把便携更新程序放入软件目录并运行更新程序.bat；旧版只识别 EXE 同级的 update.ps1。"),
        ("更新后账号或任务消失", "data 目录是否被移动或删除。", "停止继续操作，检查原软件目录的 data 和备份；正常升级不会覆盖 data。"),
    ], widths=[4.3, 6.1, 6.2])
    add_text(doc, "排查问题时请保留任务编号、账号编号、平台、发生时间和后台日志原文。不要在未记录日志的情况下反复点击真实发送或删除任务，这会增加复现难度。")

    doc.add_page_break()
    add_heading(doc, "附录 操作安全边界", 1)
    add_bullet(doc, "模拟回复、模拟私信和浏览器诊断只允许填入或读取，不点击最终发送。")
    add_bullet(doc, "真实发送和真实发布开关默认关闭，开启前必须确认账号、平台、目标用户、原评论和完整正文。")
    add_bullet(doc, "任务暂停、人工验证、失败和无更多状态都必须以任务中心和日志记录为准，不以单个页面画面猜测。")
    add_bullet(doc, "升级只替换程序，不替换 data、LLM 配置和日志；操作前不要手动删除这些目录。")
    add_bullet(doc, "员工只操作自己的账号和数据；管理员审批、查看全员数据和执行恢复操作。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
